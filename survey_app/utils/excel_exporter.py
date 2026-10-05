from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models.template import Template
from models.session import Session


def export_session_to_excel(session: Session, template: Template, output_path: str):
    """
    يصدّر نتائج جلسة التفريغ إلى ملف Excel يحتوي على ورقتي عمل:
    1. النتائج النصية: النصوص الأصلية للإجابات.
    2. النتائج الرقمية: القيم الرقمية (الترتيب) للإجابات ذات الخيارات، والنصوص للأسئلة العامة.
    """
    wb = openpyxl.Workbook()
    
    # ─── الورقة الأولى: النتائج النصية ──────────────────────────────────────────
    ws_text = wb.active
    ws_text.title = "النتائج النصية"
    _format_sheet(ws_text, session, template, mode="text")

    # ─── الورقة الثانية: النتائج الرقمية ────────────────────────────────────────
    ws_numeric = wb.create_sheet("النتائج الرقمية")
    _format_sheet(ws_numeric, session, template, mode="numeric")

    wb.save(output_path)


def _format_sheet(ws, session: Session, template: Template, mode: str):
    """تنسيق وتعبئة ورقة عمل محددة"""
    ws.sheet_view.rightToLeft = True

    # ألوان وأنماط
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    alt_fill = PatternFill("solid", fgColor="D6E4F0")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    # ─── رأس الجدول ───────────────────────────────────────────────────────────
    ws.cell(row=1, column=1, value="رقم الاستمارة")
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=1).font = header_font
    ws.cell(row=1, column=1).alignment = center_align
    ws.cell(row=1, column=1).border = thin_border
    ws.column_dimensions["A"].width = 16

    questions = sorted(template.questions, key=lambda q: q.column_index)
    for col_offset, question in enumerate(questions, start=2):
        cell = ws.cell(row=1, column=col_offset, value=question.text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        col_letter = get_column_letter(col_offset)
        ws.column_dimensions[col_letter].width = max(20, len(question.text) + 4)

    # ─── بيانات الاستمارات ────────────────────────────────────────────────────
    from models.question import QuestionType

    for form in session.forms:
        row_num = form.form_index + 2
        use_alt = form.form_index % 2 == 1

        # رقم الاستمارة
        num_cell = ws.cell(row=row_num, column=1, value=form.form_index + 1)
        num_cell.alignment = center_align
        num_cell.border = thin_border
        if use_alt:
            num_cell.fill = alt_fill

        # إجابات كل سؤال
        for col_offset, question in enumerate(questions, start=2):
            raw_answer = form.answers.get(question.id, "")
            
            display_value = raw_answer
            if mode == "numeric" and raw_answer:
                if question.question_type in (QuestionType.DEMOGRAPHIC_SINGLE, QuestionType.LIKERT, QuestionType.DEMOGRAPHIC_SINGLE_OTHER):
                    # إيجاد ترتيب الإجابة (1-based)
                    try:
                        idx = question.answers.index(raw_answer)
                        display_value = idx + 1
                    except ValueError:
                        display_value = raw_answer
                elif question.question_type == QuestionType.DEMOGRAPHIC_MULTIPLE:
                    # تحويل الإجابات المتعددة إلى أرقام مفصولة بفواصل
                    selected = [s.strip() for s in raw_answer.split(",") if s.strip()]
                    indices = []
                    for s in selected:
                        try:
                            indices.append(str(question.answers.index(s) + 1))
                        except ValueError:
                            indices.append(s)
                    display_value = ",".join(indices)

            cell = ws.cell(row=row_num, column=col_offset, value=display_value)
            cell.alignment = right_align if isinstance(display_value, str) else center_align
            cell.border = thin_border
            if use_alt:
                cell.fill = alt_fill

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
